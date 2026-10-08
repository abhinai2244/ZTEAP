from pathlib import Path
import math
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "deliverables"
ASSET_DIR = OUT_DIR / "srs_assets"
OUT_DIR.mkdir(exist_ok=True)
ASSET_DIR.mkdir(exist_ok=True)

NAVY = "17365D"
BLUE = "2F75B5"
PALE = "EAF2F8"
LIGHT = "F4F6F7"
GRID = "D9E2F3"
TEXT = "1F2937"


FONT_REG = Path("C:/Windows/Fonts/arial.ttf")
FONT_BOLD = Path("C:/Windows/Fonts/arialbd.ttf")


def font(size, bold=False):
    return ImageFont.truetype(str(FONT_BOLD if bold else FONT_REG), size)


def canvas(title, size=(2100, 1250)):
    im = Image.new("RGB", size, "white")
    d = ImageDraw.Draw(im)
    d.text((50, 35), title, font=font(42, True), fill="#17365D")
    return im, d


def px(x, y, size):
    return int(x * size[0]), int((1 - y) * size[1])


def draw_box(d, size, xy, w, h, text, fc="#EAF2F8", ec="#2F75B5", fontsize=24):
    x, y = xy
    left = int(x * size[0]); top = int((1-y-h) * size[1]); right = int((x+w)*size[0]); bottom = int((1-y)*size[1])
    d.rounded_rectangle((left,top,right,bottom), radius=18, fill=fc, outline=ec, width=4)
    f = font(fontsize, True)
    lines = text.split("\n")
    heights = [d.textbbox((0,0), line, font=f)[3] for line in lines]
    total = sum(heights) + max(0,len(lines)-1)*4
    cy = (top+bottom-total)//2
    for line, lh in zip(lines, heights):
        bb=d.textbbox((0,0),line,font=f); tw=bb[2]-bb[0]
        d.text(((left+right-tw)//2,cy),line,font=f,fill="#172B4D")
        cy += lh+4


def arrow(d, size, start, end, label=None, dashed=False):
    x1,y1=px(*start,size); x2,y2=px(*end,size)
    if dashed:
        for i in range(0,20,2):
            a=i/20; b=min((i+1)/20,1)
            d.line((x1+(x2-x1)*a,y1+(y2-y1)*a,x1+(x2-x1)*b,y1+(y2-y1)*b),fill="#5E6C84",width=4)
    else:
        d.line((x1,y1,x2,y2),fill="#5E6C84",width=4)
    ang=math.atan2(y2-y1,x2-x1); al=18
    for delta in (.55,-.55):
        d.line((x2,y2,x2-al*math.cos(ang+delta),y2-al*math.sin(ang+delta)),fill="#5E6C84",width=4)
    if label:
        mx,my=(x1+x2)//2,(y1+y2)//2
        f=font(18); bb=d.textbbox((0,0),label,font=f)
        d.rectangle((mx-4,my-24,mx+(bb[2]-bb[0])+6,my+3),fill="white")
        d.text((mx,my-22),label,font=f,fill="#42526E")


def make_architecture(path):
    im, d = canvas("Zero Trust Access Portal Architecture")
    size=im.size
    layers = [
        (.09, .77, .82, .11, "Client Layer\nBrowser and role-specific dashboards", "#E8F1FB"),
        (.09, .61, .82, .11, "Application Layer\nNext.js App Router, route handlers, validation and session middleware", "#EAF7F1"),
        (.09, .41, .82, .15, "Zero Trust Security Core\nAuthentication  |  Server-side RBAC  |  Policy Engine  |  Risk Engine  |  Approval Service  |  Audit Logger", "#FFF4E5"),
        (.09, .25, .82, .11, "Domain and Data Access Layer\nUser, Device, Resource, Policy, Access Request and Approval services  |  Prisma ORM", "#F4ECF7"),
        (.25, .08, .50, .10, "PostgreSQL 16\nRelational records, decisions, assessments and audit events", "#FDEDEC"),
    ]
    for x, y, w, h, label, fc in layers:
        draw_box(d,size,(x,y),w,h,label,fc=fc,fontsize=23)
    for y1, y2 in [(.77, .72), (.61, .56), (.41, .36), (.25, .18)]:
        arrow(d,size,(.5,y1),(.5,y2))
    d.text((1900,430),"VERIFY\nEVERY\nREQUEST",font=font(22,True),fill="#B54708",anchor="mm",align="center")
    im.save(path)


def make_use_case(path):
    im,d=canvas("Use Case Model",(2100,1350)); size=im.size
    actors = {"Employee": (.07, .72), "Security\nAdministrator": (.07, .34),
              "Resource\nOwner": (.91, .70), "Auditor": (.91, .30)}
    for name, (x, y) in actors.items():
        hx,hy=px(x,y+.06,size); d.ellipse((hx-24,hy-24,hx+24,hy+24),fill="white",outline="#17365D",width=4)
        _,by=px(x,y-.025,size); d.line((hx,hy+24,hx,by),fill="#17365D",width=4)
        lx,ly=px(x-.03,y+.01,size); rx,ry=px(x+.03,y+.01,size); d.line((lx,ly,rx,ry),fill="#17365D",width=4)
        llx,lly=px(x-.03,y-.07,size); rrx,rry=px(x+.03,y-.07,size); d.line((hx,by,llx,lly),fill="#17365D",width=4); d.line((hx,by,rrx,rry),fill="#17365D",width=4)
        tx,ty=px(x,y-.10,size); d.multiline_text((tx,ty),name,font=font(21,True),fill="#172B4D",anchor="ma",align="center")
    cases = [
        ("Authenticate and\nend session", .33, .79), ("Request access", .33, .62),
        ("Inspect decision\nand risk factors", .33, .45), ("Manage users,\nroles and devices", .33, .28),
        ("Manage resources\nand policies", .67, .72), ("Approve or reject\nrequests", .67, .53),
        ("Review audit and\nsecurity events", .67, .32), ("Export SIEM data", .67, .15),
        ("Evaluate policy and\ncalculate risk", .50, .50),
    ]
    for label, x, y in cases:
        draw_box(d,size,(x-.105,y-.045),.21,.09,label,fc="#F7F9FC",fontsize=20)
    for a,b in [((.10,.72),(.225,.79)),((.10,.72),(.225,.62)),((.10,.72),(.225,.45)),
                ((.10,.34),(.225,.28)),((.10,.34),(.565,.32)),((.88,.70),(.775,.72)),
                ((.88,.70),(.775,.53)),((.88,.30),(.775,.32)),((.88,.30),(.775,.15))]: arrow(d,size,a,b)
    arrow(d,size,(.435,.62),(.475,.53),"include",True)
    l,t=px(.20,.88,size); r,b=px(.80,.08,size); d.rounded_rectangle((l,t,r,b),radius=20,outline="#A5ADBA",width=3)
    d.text(px(.22,.105,size),"ZTAP system boundary",font=font(18),fill="#6B778C")
    im.save(path)


def make_dfd(path):
    im,d=canvas("Data Flow Diagram Level 1",(2200,1350)); size=im.size
    entities = [("Employee",.03,.70),("Administrator",.03,.22),("Resource Owner",.84,.68),("Auditor",.84,.20)]
    for t,x,y in entities: draw_box(d,size,(x,y),.13,.09,t,fc="#FFFFFF",ec="#42526E",fontsize=20)
    procs=[("P1\nAuthentication",.23,.74),("P2\nGovernance",.23,.24),("P3\nAccess Request",.43,.64),
           ("P4\nPolicy and Risk",.43,.39),("P5\nApproval",.64,.61),("P6\nAudit",.64,.25)]
    for t,x,y in procs: draw_box(d,size,(x,y),.15,.10,t,fc="#EAF2F8",fontsize=20)
    stores=[("D1 Users and\nSessions",.22,.51),("D2 Resources,\nPolicies, Devices",.22,.08),
            ("D3 Requests, Decisions\nand Assessments",.43,.08),("D4 Audit Log",.64,.08)]
    for t,x,y in stores: draw_box(d,size,(x,y),.16,.075,t,fc="#FFF8E7",ec="#B54708",fontsize=17)
    flows=[((.16,.745),(.23,.79),"credentials"),((.305,.74),(.30,.585),"verify/session"),
           ((.16,.735),(.43,.69),"request"),((.505,.64),(.505,.49),"context"),
           ((.43,.44),(.38,.545),"identity/device"),((.58,.44),(.64,.66),"approval needed"),
           ((.79,.66),(.84,.72),"review"),((.71,.61),(.55,.18),"outcome"),
           ((.16,.265),(.23,.29),"manage"),((.305,.24),(.30,.155),"updates"),
           ((.84,.245),(.79,.30),"query"),((.715,.25),(.72,.155),"read"),
           ((.58,.42),(.64,.30),"event"),((.505,.39),(.51,.155),"decision"),
           ((.30,.51),(.43,.44),"attributes")]
    for s,e,l in flows: arrow(d,size,s,e,l)
    im.save(path)


def make_er(path):
    im,d=canvas("Entity Relationship Diagram",(2200,1350)); size=im.size
    nodes={
        "User":(.06,.69),"Role":(.06,.43),"Permission":(.06,.18),"UserRole":(.25,.53),
        "Device":(.25,.76),"Session":(.25,.27),"Resource":(.45,.76),"ResourcePolicy":(.67,.76),
        "AccessRequest":(.45,.48),"RiskAssessment":(.67,.52),"AccessDecision":(.67,.35),
        "Approval":(.84,.48),"AuditLog":(.45,.17)}
    subtitles={"User":"identity and state","Role":"named authority","Permission":"action and resource",
               "UserRole":"user role link","Device":"posture and trust","Session":"expiry and token",
               "Resource":"owner and sensitivity","ResourcePolicy":"time role risk rules",
               "AccessRequest":"user resource device","RiskAssessment":"score level factors",
               "AccessDecision":"allow deny step up","Approval":"approver action reason","AuditLog":"append only event"}
    for n,(x,y) in nodes.items(): draw_box(d,size,(x,y),.135,.085,n+"\n"+subtitles[n],fc="#F7F9FC",fontsize=16)
    rels=[("User","UserRole","1 : many"),("Role","UserRole","1 : many"),("Role","Permission","1 : many"),
          ("User","Device","1 : many"),("User","Session","1 : many"),("User","AccessRequest","1 : many"),
          ("User","Resource","owns"),("Resource","ResourcePolicy","1 : many"),("Resource","AccessRequest","1 : many"),
          ("Device","AccessRequest","1 : many"),("AccessRequest","RiskAssessment","1 : 1"),
          ("AccessRequest","AccessDecision","1 : 1"),("AccessRequest","Approval","1 : many"),
          ("User","AuditLog","actor")]
    for a,b,l in rels:
        xa,ya=nodes[a]; xb,yb=nodes[b]
        arrow(d,size,(xa+.135,ya+.042),(xb,yb+.042),l)
    im.save(path)


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=110, bottom=110, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar"); tcPr.append(tcMar)
    for m, v in (("top",top),("start",start),("bottom",bottom),("end",end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node=OxmlElement(f"w:{m}"); tcMar.append(node)
        node.set(qn("w:w"),str(v)); node.set(qn("w:type"),"dxa")


def add_table(doc, headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0].cells
    for i,h in enumerate(headers):
        hdr[i].text=h; set_cell_shading(hdr[i],NAVY)
        hdr[i].vertical_alignment=WD_ALIGN_VERTICAL.CENTER
        for r in hdr[i].paragraphs[0].runs:
            r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(font_size)
    for ri,row in enumerate(rows):
        cells=table.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=str(val); cells[i].vertical_alignment=WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cells[i])
            if ri%2: set_cell_shading(cells[i],"F4F7FB")
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after=Pt(0)
                for run in p.runs: run.font.size=Pt(font_size); run.font.color.rgb=RGBColor.from_string(TEXT)
    if widths:
        for row in table.rows:
            for i,w in enumerate(widths): row.cells[i].width=Cm(w)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    return table


def add_bullets(doc, items, level=0):
    for item in items:
        p=doc.add_paragraph(style="List Bullet" if level==0 else "List Bullet 2")
        p.add_run(item)


def add_caption(doc, text):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(8)
    r=p.add_run(text); r.italic=True; r.font.size=Pt(8.5); r.font.color.rgb=RGBColor(89,89,89)


def add_picture(doc, path, caption, width=6.7):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def keep_with_next(paragraph):
    pPr=paragraph._p.get_or_add_pPr(); el=OxmlElement("w:keepNext"); pPr.append(el)


def build_doc():
    arch=ASSET_DIR/"architecture.png"; uc=ASSET_DIR/"use_case.png"; dfd=ASSET_DIR/"dfd_level_1.png"; er=ASSET_DIR/"er_diagram.png"
    make_architecture(arch); make_use_case(uc); make_dfd(dfd); make_er(er)

    doc=Document()
    sec=doc.sections[0]; sec.top_margin=Cm(1.8); sec.bottom_margin=Cm(1.7); sec.left_margin=Cm(2.0); sec.right_margin=Cm(2.0)
    styles=doc.styles
    normal=styles["Normal"]; normal.font.name="Aptos"; normal.font.size=Pt(10); normal.font.color.rgb=RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.08
    for style_name,size in [("Title",28),("Subtitle",12),("Heading 1",17),("Heading 2",13),("Heading 3",11)]:
        st=styles[style_name]; st.font.name="Aptos Display" if style_name!="Normal" else "Aptos"; st.font.color.rgb=RGBColor(0,0,0); st.font.size=Pt(size)
        if style_name.startswith("Heading"): st.font.bold=True; st.paragraph_format.space_before=Pt(12); st.paragraph_format.space_after=Pt(5); st.paragraph_format.keep_with_next=True

    # Header and footer
    header=sec.header.paragraphs[0]; header.text="ZERO TRUST ACCESS PORTAL  |  SOFTWARE REQUIREMENTS SPECIFICATION"
    header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    for r in header.runs: r.font.size=Pt(7.5); r.font.color.rgb=RGBColor.from_string("6B7280")
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=footer.add_run("ZTAP SRS  |  Version 1.0  |  8 October 2026")
    r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string("6B7280")

    # Cover
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(58); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("Software Requirements Specification"); r.bold=True; r.font.size=Pt(29); r.font.color.rgb=RGBColor(0,0,0)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("Zero Trust Enterprise Access Portal"); r.bold=True; r.font.size=Pt(20); r.font.color.rgb=RGBColor.from_string(NAVY)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("Risk adaptive identity device and policy based enterprise access control"); r.font.size=Pt(11); r.font.color.rgb=RGBColor.from_string("4B5563")
    doc.add_paragraph()
    add_table(doc,["Document control","Value"],[
        ("Document ID","ZTAP SRS 001"),("Version","1.0"),("Status","Baseline for implementation and Scrum planning"),
        ("Prepared for","Academic software engineering evaluation"),("Prepared on","8 October 2026"),
        ("System baseline","Next.js 14 TypeScript Prisma PostgreSQL implementation in the project repository")
    ],[4.2,11.5],9)
    doc.add_paragraph()
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("Purpose"); r.bold=True; r.font.size=Pt(12)
    p=doc.add_paragraph("This specification defines the behavior, security controls, interfaces, data model, quality attributes and acceptance basis for the Zero Trust Enterprise Access Portal. It is the authoritative source for backlog refinement, verification and academic assessment.")
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    doc.add_heading("Document control",1)
    add_table(doc,["Version","Date","Change","Owner"],[
        ("1.0","8 Oct 2026","Initial repository-aligned SRS baseline","Project team")
    ],[2.2,3.0,8.0,3.0])
    doc.add_heading("Contents",1)
    toc=["1 Introduction","2 Overall description","3 Stakeholders and user classes","4 Functional requirements","5 External interface requirements","6 Data requirements","7 Architecture and design views","8 Nonfunctional requirements","9 Security and privacy requirements","10 Verification and acceptance","11 Requirements traceability","12 Constraints assumptions and risks","13 Glossary"]
    add_bullets(doc,toc)

    doc.add_heading("1 Introduction",1)
    doc.add_heading("1.1 Purpose",2)
    doc.add_paragraph("ZTAP controls access to registered enterprise applications by evaluating identity, server-side roles, device posture, resource sensitivity, policy constraints, time context and an explainable risk score on every request. The system records the resulting decision and security-relevant activity for audit and review.")
    doc.add_heading("1.2 Scope",2)
    add_bullets(doc,[
        "Authenticate users and maintain protected server-side sessions.",
        "Apply least-privilege role-based authorization at API boundaries.",
        "Register users, roles, permissions, devices, protected resources and resource policies.",
        "Evaluate access requests using deterministic policy and explainable risk engines.",
        "Route approval-required requests to resource owners and prohibit self-approval.",
        "Provide employee, administrator, resource-owner and auditor dashboards.",
        "Maintain an append-only audit trail and provide structured SIEM-oriented export.",
        "Support containerized deployment and automated unit, integration, fuzz and end-to-end testing."
    ])
    doc.add_heading("1.3 Out of scope",2)
    doc.add_paragraph("The baseline does not provide a production identity-provider federation, hardware-backed device attestation, a full SIEM product, automated provisioning into third-party applications, or a mobile client. These are future integrations, not implied current behavior.")
    doc.add_heading("1.4 Product objective",2)
    doc.add_paragraph("The primary objective is to demonstrate that access is never granted only because a user reached the portal. Each access request must be independently authenticated, authorized, risk assessed, policy evaluated, decided and logged.")

    doc.add_heading("2 Overall description",1)
    doc.add_heading("2.1 Product perspective",2)
    doc.add_paragraph("ZTAP is a full-stack web application built with the Next.js App Router and TypeScript. Browser dashboards call protected route handlers. Middleware validates an encrypted session and roles before domain services access PostgreSQL through Prisma. The policy and risk engines return deterministic decisions and evidence used by access-request, approval and audit services.")
    doc.add_heading("2.2 Operating environment",2)
    add_table(doc,["Area","Baseline"],[
        ("Client","Modern browser with JavaScript and secure cookie support"),("Application runtime","Node.js 20 with Next.js 14"),
        ("Database","PostgreSQL 16 accessed through Prisma ORM"),("Local deployment","npm or Docker Compose"),
        ("Orchestration","Kubernetes manifests for the ztap-system namespace"),("Testing","Vitest and Playwright")
    ],[4.0,11.5])
    doc.add_heading("2.3 Product capabilities",2)
    add_table(doc,["Capability","Implemented behavior"],[
        ("Identity","Email and password login, Argon2id verification, lockout tracking, encrypted sessions and logout"),
        ("Governance","User creation, account disablement, role assignment and revocation with anti-self-escalation protection"),
        ("Device trust","Device registration and administrator-controlled TRUSTED, COMPLIANT, UNTRUSTED or COMPROMISED posture"),
        ("Protected resources","Resource ownership, sensitivity, required role, risk level, active state and approval requirement"),
        ("Decisioning","Identity, role, device, resource, temporal and risk evaluation producing ALLOW, DENY, STEP_UP_AUTHENTICATION or REQUIRE_APPROVAL"),
        ("Audit","Severity, actor, action, result, correlation ID and scrubbed metadata stored as security events")
    ],[4.0,11.5])

    doc.add_heading("3 Stakeholders and user classes",1)
    add_table(doc,["User class","Goals","Key restrictions"],[
        ("Employee","Sign in, register or inspect devices, request access and understand decisions","Can view only permitted personal data and cannot bypass policy"),
        ("Resource Owner","Manage owned applications and decide pending approvals","Cannot approve own access request"),
        ("Security Administrator","Manage users, roles, device posture, policies and security status","Cannot elevate own roles through the protected role-management flow"),
        ("Auditor","Inspect events, decisions and exported evidence","Read-only access to operational data"),
        ("System Administrator","Perform highest-level governance operations","All actions remain auditable")
    ],[3.2,7.0,5.3])

    doc.add_heading("4 Functional requirements",1)
    frs=[
        ("FR-01","User authentication","The system shall authenticate an active user by email and Argon2id-verified password and create a protected session after valid credentials."),
        ("FR-02","Login abuse control","The system shall record failed login attempts, enforce temporary account lockout after the configured threshold and rate-limit repeated authentication attempts."),
        ("FR-03","Session termination","The system shall invalidate the user session on logout and prevent reuse of the terminated session."),
        ("FR-04","User lifecycle","Authorized administrators shall list users, create user accounts and disable an account immediately."),
        ("FR-05","Role governance","Authorized administrators shall assign and revoke roles; the system shall prevent protected self-privilege escalation."),
        ("FR-06","Server-side authorization","Every protected API operation shall validate the authenticated session and required role or permission on the server."),
        ("FR-07","Device registration","An authenticated user shall register a device with name and optional operating-system and browser attributes."),
        ("FR-08","Device posture administration","Authorized administrators shall update device status, managed state, compliance state and trust score."),
        ("FR-09","Resource registry","Resource owners or administrators shall create and manage internal application records including owner, sensitivity, required role, risk level and approval flag."),
        ("FR-10","Resource policy management","Authorized users shall create, update and delete policies containing role, maximum risk, time, day and managed-device constraints."),
        ("FR-11","Access request submission","An employee shall submit a reasoned access request for a selected resource and owned device."),
        ("FR-12","Context validation","The system shall reject requests involving inactive users or resources, unknown entities, unauthorized devices or prohibited role combinations."),
        ("FR-13","Risk assessment","The system shall calculate a score from 0 to 100, classify its level and persist human-readable contributing factors."),
        ("FR-14","Dynamic policy evaluation","The system shall evaluate identity, role, device, resource, temporal and risk constraints for every access request."),
        ("FR-15","Decision persistence","The system shall persist exactly one decision and one risk assessment for each evaluated access request."),
        ("FR-16","Approval routing","If policy or resource configuration requires approval, the system shall place the request in a pending state for its resource owner."),
        ("FR-17","Approval decision","An authorized resource owner shall approve, reject or request additional authentication with a recorded reason."),
        ("FR-18","Segregation of duties","The system shall reject and audit any attempt by a requester to approve their own request."),
        ("FR-19","Audit inspection","An auditor or authorized administrator shall filter and inspect security events with actor, action, result, severity, time and correlation data."),
        ("FR-20","Operational dashboards","The system shall present role-appropriate summaries for employees, resource owners, administrators and auditors.")
    ]
    add_table(doc,["ID","Name","Requirement"],frs,[1.6,3.8,10.1],8.2)

    doc.add_heading("5 External interface requirements",1)
    doc.add_heading("5.1 User interface",2)
    doc.add_paragraph("The interface shall provide a responsive login page and separate employee, owner, administrator and auditor dashboards. Security decisions shall use unambiguous text and status styling and shall expose the decision reason and risk factors without disclosing secrets.")
    doc.add_heading("5.2 API interface",2)
    api_rows=[
        ("/api/auth/login and /logout","Create or terminate an authenticated session"),
        ("/api/auth/me","Return the current safe user context"),
        ("/api/users and /api/users/{id}","List, create, inspect or disable users"),
        ("/api/users/{id}/roles","Assign or revoke roles"),
        ("/api/devices and /api/devices/{id}","Register, list and administer devices"),
        ("/api/resources and /api/resources/{id}","Register, list and manage protected applications"),
        ("/api/policies and /api/policies/{id}","Create, modify or delete resource policies"),
        ("/api/access/requests","Submit and list access requests"),
        ("/api/approvals and /api/approvals/{id}","List and decide approval work"),
        ("/api/audit","Query audit evidence"),
        ("/api/stats/admin","Return security-administration metrics")
    ]
    add_table(doc,["Endpoint group","Responsibility"],api_rows,[6.0,9.5])
    doc.add_heading("5.3 Communication and validation",2)
    doc.add_paragraph("Client and API communication shall use HTTPS in non-local environments. JSON request bodies shall be parsed safely and validated with explicit runtime schemas. Malformed input shall return a client error without stack traces or secret values.")

    doc.add_heading("6 Data requirements",1)
    doc.add_heading("6.1 Logical entities",2)
    add_table(doc,["Entity","Purpose","Important integrity rule"],[
        ("User","Identity profile and account state","Email is unique and password hash is never returned"),
        ("Role Permission UserRole","RBAC governance","A user-role pair and role permission are unique"),
        ("Device","Endpoint posture","Each device belongs to one user"),
        ("Resource ResourcePolicy","Protected application and constraints","Resource name is unique and policies cascade with resource deletion"),
        ("AccessRequest","Requested access context","References one user resource and device"),
        ("RiskAssessment AccessDecision","Explainability and outcome","Exactly one of each per evaluated request"),
        ("Approval","Human decision record","References request and approver"),
        ("Session","Server-verifiable authenticated context","Token is unique and expiration is required"),
        ("AuditLog","Append-only security evidence","Indexed by time actor action severity and correlation ID")
    ],[4.2,6.5,5.0])
    doc.add_heading("6.2 Data retention and privacy",2)
    doc.add_paragraph("Production retention periods shall be configured by organizational policy. Password hashes, session tokens and secrets shall not appear in dashboards, logs or exports. Audit metadata shall be scrubbed before persistence. Access to personal and security data shall follow least privilege.")
    add_picture(doc,er,"Figure 1 Entity relationship model aligned with the Prisma schema")

    doc.add_heading("7 Architecture and design views",1)
    doc.add_heading("7.1 Layered architecture",2)
    doc.add_paragraph("The design separates presentation, transport and authorization, security decisioning, domain services and persistence. This separation makes policy rules testable and prevents the browser from becoming an authority for access decisions.")
    add_picture(doc,arch,"Figure 2 Layered system architecture")
    doc.add_heading("7.2 Use case view",2)
    add_picture(doc,uc,"Figure 3 Primary actors and use cases")
    doc.add_heading("7.3 Data flow view",2)
    add_picture(doc,dfd,"Figure 4 Level 1 data flow across trust boundaries")
    doc.add_heading("7.4 Decision sequence",2)
    add_bullets(doc,[
        "Authenticate the session and load the safe user context.",
        "Validate the request, resource and device ownership context.",
        "Calculate and persist the explainable risk assessment.",
        "Evaluate identity, role, device, resource, temporal and risk policies.",
        "Persist the decision and update request status.",
        "Create an approval path when required and write correlated audit events.",
        "Return the minimum decision evidence needed by the requesting dashboard."
    ])

    doc.add_heading("8 Nonfunctional requirements",1)
    nfr=[
        ("NFR-01","Security","Passwords shall use Argon2id; sessions shall use HttpOnly, Secure in production and SameSite protections."),
        ("NFR-02","Performance","For normal academic demonstration load, interactive API operations should complete within 2 seconds excluding deployment cold start."),
        ("NFR-03","Availability","Container health and restart behavior shall allow the application to recover from process failure without corrupting persisted data."),
        ("NFR-04","Reliability","Access decisions and risk assessments shall be transactionally consistent with their access request."),
        ("NFR-05","Auditability","Security-sensitive actions shall produce timestamped, correlated and severity-classified audit events."),
        ("NFR-06","Maintainability","Security engines and domain services shall remain modular, typed and covered by automated tests."),
        ("NFR-07","Usability","Role dashboards shall be responsive, keyboard operable and show clear success, error and empty states."),
        ("NFR-08","Portability","The application shall run locally and in Docker; Kubernetes manifests shall use non-root execution and resource constraints."),
        ("NFR-09","Scalability","List endpoints shall support pagination and database indexes shall cover frequent identity, status and audit queries."),
        ("NFR-10","Observability","Errors and security events shall be diagnosable through structured application and audit data without leaking secrets.")
    ]
    add_table(doc,["ID","Quality","Requirement"],nfr,[1.7,3.0,10.8],8.3)

    doc.add_heading("9 Security and privacy requirements",1)
    add_table(doc,["Control area","Required control"],[
        ("Authentication","Credential verification, failed-attempt handling, lockout and secure session lifecycle"),
        ("Authorization","Default-deny server-side checks and no trust in client-supplied role state"),
        ("Input security","Schema validation, parameterized ORM access and safe malformed-JSON handling"),
        ("Separation of duties","Block self-approval and protected self-role modification"),
        ("Device security","Compromised devices cause denial; managed-device rules are policy enforceable"),
        ("Data protection","No password hash or session secret in API payloads; secrets supplied through environment configuration"),
        ("Audit protection","No user-facing mutation route for audit records; metadata is scrubbed"),
        ("Deployment","Non-root container, dropped Linux capabilities, resource limits and isolated database network"),
        ("Security headers","Apply browser security headers and production HTTPS controls")
    ],[4.2,11.5])

    doc.add_heading("10 Verification and acceptance",1)
    add_table(doc,["Test level","Evidence","Acceptance condition"],[
        ("Unit","Password, risk and policy engine tests","All deterministic rules and boundary cases pass"),
        ("Integration","RBAC and privilege-escalation tests","Unauthorized operations are rejected and audited"),
        ("Fuzz","Malicious and malformed input suite","Application handles payloads without unsafe execution or secret disclosure"),
        ("End to end","Browser authentication and role journeys","Critical user journeys complete with expected decisions"),
        ("Build","Production compilation","Type checking and optimized application build succeed"),
        ("Demonstration","Role-based demonstration script","Evaluator can reproduce allow deny approval and audit scenarios")
    ],[3.0,5.5,7.2])
    doc.add_heading("10.1 Definition of done",2)
    add_bullets(doc,[
        "Acceptance criteria pass and code is reviewed.",
        "Authorization and negative-path tests exist for protected behavior.",
        "No known critical or high-severity defect remains open.",
        "Relevant audit event and error behavior are verified.",
        "Documentation and backlog traceability are updated.",
        "The item is deployed or demonstrable in the agreed environment."
    ])

    doc.add_heading("11 Requirements traceability",1)
    trace=[
        ("FR-01 FR-02 FR-03","Authentication","Auth routes password and session modules","US-01 US-02","Unit and E2E"),
        ("FR-04 FR-05 FR-06","Governance and RBAC","User service authorization middleware","US-03 US-04","Integration"),
        ("FR-07 FR-08","Device trust","Device service and device routes","US-05 US-15","Unit and integration"),
        ("FR-09 FR-10","Resource and policy","Resource and policy services","US-06 US-14","API integration"),
        ("FR-11 to FR-16","Access decisioning","Access request policy and risk engines","US-07 US-08 US-09","Unit integration E2E"),
        ("FR-17 FR-18","Approval workflow","Approval service and approval routes","US-10 US-16","Integration"),
        ("FR-19","Audit","AuditLogger and audit route","US-11 US-20","Integration and export"),
        ("FR-20","Dashboards","Role dashboard pages","US-12 US-13","E2E")
    ]
    add_table(doc,["Requirement","Capability","Implementation area","Backlog","Verification"],trace,[2.8,3.2,4.2,2.5,2.8],7.8)

    doc.add_heading("12 Constraints assumptions and risks",1)
    add_table(doc,["Type","Statement","Response"],[
        ("Constraint","The baseline uses a single application and PostgreSQL database.","Preserve modular services so external policy or identity systems can be added later."),
        ("Assumption","A resource has one accountable owner and users select a registered device.","Validate ownership and require administrative correction of inconsistent data."),
        ("Assumption","Server time is trustworthy for temporal policies.","Use synchronized infrastructure time in deployed environments."),
        ("Risk","Academic seed credentials may be mistaken for production credentials.","Label them as demonstration-only and replace all secrets before deployment."),
        ("Risk","Simulated posture is weaker than hardware attestation.","Treat device posture as demonstrative and document future MDM or attestation integration."),
        ("Risk","An append-only application API does not alone prevent direct database modification.","Use restricted database roles and external log forwarding in production.")
    ],[2.0,7.0,6.5],8.2)

    doc.add_heading("13 Glossary",1)
    add_table(doc,["Term","Meaning"],[
        ("Zero Trust","A security model that continuously verifies each request instead of trusting network location."),
        ("RBAC","Role-based access control that maps roles to permitted operations."),
        ("Risk score","A bounded numeric assessment with stored contributing factors."),
        ("Step-up authentication","A decision requiring stronger verification before access can continue."),
        ("Resource","An internal application or protected service registered in the portal."),
        ("Policy","Constraints governing role device time day and acceptable risk for a resource."),
        ("SIEM","Security information and event management platform consuming structured security events."),
        ("Correlation ID","Identifier linking related operations and audit evidence across a request flow.")
    ],[4.0,11.7])

    doc.add_heading("Approval",1)
    doc.add_paragraph("This version is suitable as the requirements baseline for Jira backlog reconstruction, two-sprint planning and final project reporting. Changes to scope should be reflected in both this document and the linked backlog.")
    add_table(doc,["Role","Name","Signature","Date"],[
        ("Project supervisor","","",""),("Product owner","","",""),("Development representative","","","")
    ],[4.0,4.0,4.0,3.0])

    # Global paragraph cleanup
    for p in doc.paragraphs:
        for run in p.runs:
            run.font.name = run.font.name or "Aptos"
        if p.style.name.startswith("Heading"):
            keep_with_next(p)

    out=OUT_DIR/"ZTAP_Software_Requirements_Specification.docx"
    doc.save(out)
    print(out)


if __name__ == "__main__":
    build_doc()
